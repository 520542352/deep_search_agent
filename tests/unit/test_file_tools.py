from pathlib import Path

import pandas as pd
import pytest
import docx
from pypdf import PdfReader, PdfWriter

from api.context import set_session_context
from tools import markdown_tool, pdf_tool, upload_file_read_tool
from utils import pdf_renderer


@pytest.mark.unit
def test_generate_and_read_markdown(tmp_path: Path) -> None:
    set_session_context(str(tmp_path))

    generated = markdown_tool.generate_markdown.invoke(
        {"content": "# 标题\n\n正文", "filename": "report", "path": "nested"}
    )
    content = upload_file_read_tool.read_file.invoke(
        {"filename": "nested/report.md"}
    )

    assert "已成功生成" in generated
    assert content == "# 标题\n\n正文"
    assert (tmp_path / "nested" / "report.md").exists()


@pytest.mark.unit
def test_read_file_rejects_path_outside_session(tmp_path: Path) -> None:
    set_session_context(str(tmp_path / "session"))

    result = upload_file_read_tool.read_file.invoke({"filename": "../secret.txt"})

    assert result.startswith("错误:")
    assert "拒绝访问" in result


@pytest.mark.unit
def test_read_excel_returns_preview_and_shape(tmp_path: Path) -> None:
    set_session_context(str(tmp_path))
    workbook = tmp_path / "sample.xlsx"
    pd.DataFrame({"name": ["A", "B"], "value": [1, 2]}).to_excel(
        workbook, index=False
    )

    result = upload_file_read_tool.read_file.invoke({"filename": "sample.xlsx"})

    assert "行数2,列数：2" in result
    assert "name,value" in result
    assert "[前5行数据预览]" in result


@pytest.mark.unit
def test_convert_pdf_uses_resolved_session_paths(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    set_session_context(str(tmp_path))
    (tmp_path / "report.md").write_text("# Report", encoding="utf-8")
    captured: dict[str, Path] = {}

    def fake_convert(source: Path, target: Path) -> str:
        captured["source"] = source
        captured["target"] = target
        return "converted"

    monkeypatch.setattr(pdf_tool, "convert_md_to_pdf_via_html", fake_convert)

    result = pdf_tool.convert_md_to_pdf.invoke(
        {"md_filename": "report.md", "pdf_filename": "exports/final.pdf"}
    )

    assert result == "converted"
    assert captured["source"] == tmp_path.resolve() / "report.md"
    assert captured["target"] == tmp_path.resolve() / "exports" / "final.pdf"


@pytest.mark.unit
def test_html_pdf_renderer_creates_readable_pdf(tmp_path: Path) -> None:
    markdown_file = tmp_path / "report.md"
    pdf_file = tmp_path / "report.pdf"
    markdown_file.write_text(
        """# 中文研究报告

This report verifies portable PDF rendering.

[Reference link](https://example.com/reference)

| 项目 | 数值 |
| --- | ---: |
| 样本 | 42 |

```python
print("研究主题")
```
""",
        encoding="utf-8",
    )

    result = pdf_renderer.convert_md_to_pdf_via_html(markdown_file, pdf_file)

    assert result.startswith("成功转换:")
    assert pdf_file.read_bytes().startswith(b"%PDF")
    reader = PdfReader(pdf_file)
    assert len(reader.pages) >= 1
    extracted = "\n".join(page.extract_text() or "" for page in reader.pages)
    assert "中文研究报告" in extracted
    assert "portable PDF rendering" in extracted
    assert "研究主题" in extracted
    assert "42" in extracted


@pytest.mark.unit
@pytest.mark.parametrize(
    "uri",
    ["https://example.com/image.png", "../outside.png", "/outside.png"],
)
def test_pdf_resource_resolver_rejects_external_or_escaped_resources(
    tmp_path: Path, uri: str
) -> None:
    with pytest.raises(ValueError):
        pdf_renderer._resolve_resource_uri(uri, tmp_path.resolve())


@pytest.mark.unit
def test_html_pdf_renderer_removes_partial_file_on_failure(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    markdown_file = tmp_path / "report.md"
    pdf_file = tmp_path / "report.pdf"
    markdown_file.write_text("# Report", encoding="utf-8")

    def fail_render(**_kwargs):
        raise RuntimeError("renderer unavailable")

    monkeypatch.setattr(pdf_renderer.pisa, "CreatePDF", fail_render)

    result = pdf_renderer.convert_md_to_pdf_via_html(markdown_file, pdf_file)

    assert result == "转换失败: renderer unavailable"
    assert not pdf_file.exists()


@pytest.mark.unit
def test_read_docx_extracts_paragraphs(tmp_path: Path) -> None:
    set_session_context(str(tmp_path))
    document = docx.Document()
    document.add_paragraph("第一段")
    document.add_paragraph("第二段")
    document.save(tmp_path / "sample.docx")

    result = upload_file_read_tool.read_file.invoke({"filename": "sample.docx"})

    assert result == "第一段\n第二段"


@pytest.mark.unit
def test_read_blank_pdf_returns_empty_text(tmp_path: Path) -> None:
    set_session_context(str(tmp_path))
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    with (tmp_path / "blank.pdf").open("wb") as output:
        writer.write(output)

    result = upload_file_read_tool.read_file.invoke({"filename": "blank.pdf"})

    assert result == ""


@pytest.mark.unit
@pytest.mark.parametrize(
    ("filename", "dependency", "expected"),
    [
        ("sample.docx", "docx", "未安装 py-docx"),
        ("sample.pdf", "pypdf", "未安装 pypdf"),
        ("sample.xlsx", "pd", "未安装pandas"),
    ],
)
def test_read_file_reports_missing_optional_dependency(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    filename: str,
    dependency: str,
    expected: str,
) -> None:
    set_session_context(str(tmp_path))
    (tmp_path / filename).write_bytes(b"placeholder")
    monkeypatch.setattr(upload_file_read_tool, dependency, None)

    result = upload_file_read_tool.read_file.invoke({"filename": filename})

    assert expected in result


@pytest.mark.unit
@pytest.mark.parametrize("filename", ["broken.docx", "broken.pdf"])
def test_read_file_converts_corrupt_document_error(
    tmp_path: Path,
    filename: str,
) -> None:
    set_session_context(str(tmp_path))
    (tmp_path / filename).write_bytes(b"not a document")

    result = upload_file_read_tool.read_file.invoke({"filename": filename})

    assert result.startswith("读取文件出错")


@pytest.mark.unit
def test_read_file_converts_corrupt_excel_error(tmp_path: Path) -> None:
    set_session_context(str(tmp_path))
    (tmp_path / "broken.xlsx").write_bytes(b"not an excel workbook")

    result = upload_file_read_tool.read_file.invoke({"filename": "broken.xlsx"})

    assert result.startswith("读取 Excel 失败")


@pytest.mark.unit
def test_read_file_handles_missing_and_binary_files(tmp_path: Path) -> None:
    set_session_context(str(tmp_path))
    missing = upload_file_read_tool.read_file.invoke({"filename": "missing.txt"})
    (tmp_path / "sample.bin").write_bytes(b"\xff\xfe\x00")
    binary = upload_file_read_tool.read_file.invoke({"filename": "sample.bin"})

    assert "不存在" in missing
    assert "不支持的文件格式'.bin'" in binary
